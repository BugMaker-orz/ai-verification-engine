#!/usr/bin/env python3
"""生成模拟合同文档（DOCX + PDF）用于 AI 验真引擎测试。

输出：
  samples/contract_main.docx        # 主合同（Word）
  samples/contract_supplement.pdf   # 补充协议（PDF，由 docx 转换）

内容与 samples/ 下已有 TXT 样例一致（故意包含跨文件冲突与欠缺条款）。
"""
from __future__ import annotations

import os
import subprocess
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

SAMPLES = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "samples")
SAMPLES = os.path.abspath(SAMPLES)


def set_font(run, name_cn: str = "宋体", size: int = 12, bold: bool = False):
    """设置中英文字体与字号。"""
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name_cn)


def add_para(doc, text: str, size: int = 12, bold: bool = False,
             align: WD_ALIGN_PARAGRAPH = WD_ALIGN_PARAGRAPH.LEFT,
             indent: bool = True):
    """添加段落。"""
    p = doc.add_paragraph()
    p.alignment = align
    if indent:
        p.paragraph_format.first_line_indent = Pt(size * 2)
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run(text)
    set_font(run, size=size, bold=bold)
    return p


def build_contract(title: str, blocks: list) -> Document:
    """按标题 + 段落块构建合同文档。"""
    doc = Document()
    # 标题
    add_para(doc, title, size=16, bold=True,
             align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    doc.add_paragraph()
    for kind, text in blocks:
        if kind == "section":      # 条款标题
            add_para(doc, text, size=12, bold=True, indent=False)
        elif kind == "blank":      # 空行
            doc.add_paragraph()
        else:                      # 正文
            add_para(doc, text, size=12)
    return doc


# ---------- 主合同 ----------
MAIN_BLOCKS = [
    ("text", "甲方：青岛海星科技有限公司"),
    ("text", "乙方：济南蓝海软件有限公司"),
    ("text", "签订日期：2026年3月15日"),
    ("text", "生效日期：2026年4月1日"),
    ("blank", ""),
    ("section", "第一条 项目内容"),
    ("text", "乙方为甲方开发企业管理系统一套，包含用户管理、订单管理、报表统计三个模块。"),
    ("section", "第二条 合同金额"),
    ("text", "合同金额：人民币 280,000 元（大写：贰拾捌万元整）。"),
    ("section", "第三条 付款方式"),
    ("text", "合同签订后7个工作日内支付30%预付款，项目验收合格后支付60%，剩余10%作为质保金，质保期满后支付。"),
    ("section", "第四条 合同期限"),
    ("text", "服务期限：自2026年4月1日至2026年12月31日。"),
    ("section", "第五条 违约责任"),
    ("text", "任何一方违反本合同约定，应向守约方支付合同总金额10%的违约金，并赔偿由此造成的实际损失。"),
    ("section", "第六条 争议解决"),
    ("text", "因本合同引起的争议，双方应友好协商解决；协商不成的，提交甲方所在地人民法院诉讼解决。"),
    ("section", "第七条 保密条款"),
    ("text", "双方对在合作过程中知悉的对方商业秘密负有保密义务，保密期限为合同终止后3年。"),
    ("section", "第八条 知识产权"),
    ("text", "本项目开发成果的知识产权归甲方所有，乙方享有署名权。"),
    ("section", "第九条 联系人"),
    ("text", "甲方联系人：张明，联系电话：13800138000"),
    ("text", "乙方联系人：李华，联系电话：13900139000"),
]

# ---------- 补充协议 ----------
SUPP_BLOCKS = [
    ("text", "甲方：青岛海星科技有限责任公司"),
    ("text", "乙方：济南蓝海软件有限公司"),
    ("text", "签订日期：2026年5月20日"),
    ("blank", ""),
    ("text", "鉴于双方于2026年3月15日签订了《软件开发服务合同》，现经协商一致，达成如下补充协议："),
    ("section", "第一条 变更内容"),
    ("text", "因需求变更，新增数据大屏模块，项目范围相应扩大。"),
    ("section", "第二条 金额变更"),
    ("text", "合同金额调整为：人民币 320,000 元（大写：叁拾贰万元整）。"),
    ("text", "原合同金额280,000元作废，以本补充协议为准。"),
    ("section", "第三条 付款方式"),
    ("text", "付款方式按原合同执行。"),
    ("section", "第四条 期限变更"),
    ("text", "服务期限延长至2027年3月31日。"),
    ("section", "第五条 其他"),
    ("text", "本补充协议未尽事宜，按原合同执行。"),
    ("text", "本补充协议一式两份，双方各执一份。"),
    ("blank", ""),
    ("text", "甲方（盖章）：__________________"),
    ("text", "乙方（盖章）：__________________"),
]


def main():
    # 1. 生成主合同 DOCX
    main_doc = build_contract("软件开发服务合同", MAIN_BLOCKS)
    main_docx = os.path.join(SAMPLES, "contract_main.docx")
    main_doc.save(main_docx)
    print(f"✓ 已生成：{main_docx}")

    # 2. 生成补充协议 DOCX（临时）
    supp_docx_tmp = os.path.join(SAMPLES, "_supplement_tmp.docx")
    supp_doc = build_contract("软件开发服务补充协议", SUPP_BLOCKS)
    supp_doc.save(supp_docx_tmp)
    print(f"✓ 已生成临时：{supp_docx_tmp}")

    # 3. 用 LibreOffice 转 PDF（跨平台：Windows 用 soffice，Linux/Mac 用 libreoffice）
    supp_pdf = os.path.join(SAMPLES, "contract_supplement.pdf")
    soffice_cmd = "soffice" if sys.platform.startswith("win") else "libreoffice"
    proc = subprocess.run(
        [soffice_cmd, "--headless", "--convert-to", "pdf",
         "--outdir", SAMPLES, supp_docx_tmp],
        capture_output=True, text=True, timeout=120,
    )
    if proc.returncode != 0:
        print(f"✗ PDF 转换失败（尝试命令: {soffice_cmd}）：{proc.stderr[:500]}", file=sys.stderr)
        print("提示：请确保已安装 LibreOffice 并加入 PATH，Windows 上通常为 soffice.exe", file=sys.stderr)
        sys.exit(1)
    # LibreOffice 以输入文件名命名输出，重命名为目标名
    tmp_pdf = os.path.join(SAMPLES, "_supplement_tmp.pdf")
    supp_pdf = os.path.join(SAMPLES, "contract_supplement.pdf")
    if os.path.isfile(tmp_pdf):
        os.replace(tmp_pdf, supp_pdf)
    os.remove(supp_docx_tmp)
    print(f"✓ 已生成：{supp_pdf}")


if __name__ == "__main__":
    main()
