"""文件解析器：支持 PDF / TXT / DOCX，统一输出纯文本与元信息。"""
from __future__ import annotations

import os
from dataclasses import dataclass, field
from typing import List


@dataclass
class ParsedDoc:
    """解析后的文档对象。"""
    path: str
    filename: str
    file_type: str          # pdf / txt / docx
    text: str               # 全文纯文本
    pages: List[str] = field(default_factory=list)   # 分页文本（PDF 专用）
    line_count: int = 0

    def lines(self) -> List[str]:
        return self.text.splitlines()


def _parse_pdf(path: str) -> ParsedDoc:
    import pdfplumber
    pages: List[str] = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            pages.append(page.extract_text() or "")
    text = "\n".join(pages)
    return ParsedDoc(
        path=path,
        filename=os.path.basename(path),
        file_type="pdf",
        text=text,
        pages=pages,
        line_count=len(text.splitlines()),
    )


def _parse_txt(path: str) -> ParsedDoc:
    with open(path, encoding="utf-8", errors="replace") as f:
        text = f.read()
    return ParsedDoc(
        path=path,
        filename=os.path.basename(path),
        file_type="txt",
        text=text,
        pages=[text],
        line_count=len(text.splitlines()),
    )


def _parse_docx(path: str) -> ParsedDoc:
    from docx import Document
    doc = Document(path)
    paragraphs = [p.text for p in doc.paragraphs]
    # 表格内容也提取
    for table in doc.tables:
        for row in table.rows:
            paragraphs.append(" | ".join(cell.text for cell in row.cells))
    text = "\n".join(paragraphs)
    return ParsedDoc(
        path=path,
        filename=os.path.basename(path),
        file_type="docx",
        text=text,
        pages=[text],
        line_count=len(text.splitlines()),
    )


_PARSERS = {
    ".pdf": _parse_pdf,
    ".txt": _parse_txt,
    ".md": _parse_txt,
    ".docx": _parse_docx,
}


def parse_file(path: str) -> ParsedDoc:
    """根据扩展名自动选择解析器。"""
    ext = os.path.splitext(path)[1].lower()
    parser = _PARSERS.get(ext)
    if parser is None:
        raise ValueError(f"不支持的文件类型: {ext}（支持 {', '.join(_PARSERS)}）")
    return parser(path)


def parse_files(paths: List[str]) -> List[ParsedDoc]:
    """批量解析。"""
    return [parse_file(p) for p in paths]
